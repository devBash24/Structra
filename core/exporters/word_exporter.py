from pathlib import Path

from docx import Document


class WordExporter:
    def export(
        self,
        extraction_result: dict,
        output_path: str | Path,
    ) -> Path:
        output_path = Path(output_path)

        document = Document()

        filename = extraction_result.get(
            "filename",
            "Extracted Document",
        )

        document.add_heading(
            f"Extracted from {filename}",
            level=1,
        )

        text = extraction_result.get("text", "")

        if text:
            document.add_heading(
                "Extracted Text",
                level=2,
            )

            for paragraph in text.split("\n"):
                paragraph = paragraph.strip()

                if paragraph:
                    document.add_paragraph(paragraph)

        tables = extraction_result.get(
            "tables",
            [],
        )

        if tables:
            document.add_heading(
                "Extracted Tables",
                level=2,
            )

        for extracted_table in tables:
            document.add_heading(
                f"Table {extracted_table['index']}",
                level=3,
            )

            columns = extracted_table["columns"]
            rows = extracted_table["rows"]

            word_table = document.add_table(
                rows=1,
                cols=len(columns),
            )

            word_table.style = "Table Grid"

            header_cells = word_table.rows[0].cells

            for index, column in enumerate(columns):
                header_cells[index].text = str(column)

            for row in rows:
                cells = word_table.add_row().cells

                for index, value in enumerate(row):
                    cells[index].text = str(value)

        document.save(output_path)

        return output_path