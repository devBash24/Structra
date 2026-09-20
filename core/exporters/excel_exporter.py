from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ExcelExporter:
    def export(
        self,
        extraction_result: dict,
        output_path: str | Path,
    ) -> Path:
        output_path = Path(output_path)

        workbook = Workbook()

        # Remove default worksheet
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        tables = extraction_result.get("tables", [])

        if not tables:
            sheet = workbook.create_sheet("Extracted Data")

            sheet["A1"] = "No tables were detected."
            sheet["A2"] = extraction_result.get("text", "")

        else:
            for table in tables:
                self._add_table_sheet(workbook, table)

        workbook.save(output_path)

        return output_path

    def _add_table_sheet(self, workbook, table: dict) -> None:
        index = table["index"]

        sheet = workbook.create_sheet(
            title=f"Table {index}"
        )

        columns = table["columns"]
        rows = table["rows"]

        # Header
        for column_index, value in enumerate(columns, start=1):
            cell = sheet.cell(
                row=1,
                column=column_index,
                value=value,
            )

            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7",
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        # Data
        for row_index, row in enumerate(rows, start=2):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value,
                )

        self._autosize_columns(sheet)

    def _autosize_columns(self, sheet) -> None:
        for column_cells in sheet.columns:
            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                value = cell.value

                if value is None:
                    continue

                max_length = max(
                    max_length,
                    len(str(value)),
                )

            sheet.column_dimensions[column_letter].width = min(
                max_length + 3,
                50,
            )